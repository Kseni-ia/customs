#!/usr/bin/env python3
"""Build HS code summary for V-TOP Europe s.r.o. / shipment WPXA26B0722CZ471."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SENDER = "HANGZHOU ONTIME I.T. CO., LTD., Hangzhou, Čína"

# (hs, popis_cz, [(ontime_code, qty, nw, gw, value)], priznak, poznamka)
ROWS = [
    (
        "8531 90 00 00",
        "Části a součásti elektrických zabezpečovacích (poplachových) zařízení – pevné tvrdé "
        "etikety a samoobslužné schránky AM/RF s magnetickým zámkem (AS, HD) a pevné štítky "
        "k zabezpečení brýlí (OP). Určeno k ochraně zboží proti krádeži v obchodech.",
        [
            ("AS1038 – ALARM PROTECTOR 8.2MHZ", 3_000, 333.00, 351.00, 6030.00),
            ("HD2001 – MINI SQUARE 8.2MHZ", 100_000, 730.00, 790.00, 2930.00),
            ("HD2013 – R50 DOME 8.2MHZ", 50_000, 620.00, 650.00, 1975.00),
            ("HD2036 – LIGHT WEIGHT TAG 58KHZ", 50_000, 390.00, 420.00, 2535.00),
            ("HD2086 – X40 WITH PIN 8.2MHZ", 10_000, 104.00, 110.00, 524.00),
            ("HD2105 – MICKEY PENCIL-B 175MM", 5_000, 53.50, 56.50, 400.00),
            ("HD2105 – MICKEY PENCIL-B 90MM", 10_000, 99.00, 105.00, 800.00),
            ("HD2290 – MITAG PENCIL 58KHZ", 1_000, 12.80, 13.40, 127.30),
            ("OP3825 – OPTICAL TAG 58KHZ", 20_000, 162.00, 174.00, 2666.00),
        ],
        "",
        "Sloučeno 9 položek faktury (skupiny AS + HD + OP) do jednoho HS kódu.",
    ),
    (
        "8505 11 10",
        "Permanentní magnety z kovu – magnetický oddělovač (detacher) pro pevný štítek zboží "
        "v obchodech.",
        [
            ("DT4114 – SUPER DETACHER W/ LANYARD", 30, 9.20, 9.80, 384.00),
            ("DT4088 – MITAG DETACHER", 20, 18.10, 19.30, 796.00),
            ("DT4011 – DT053 SUPER SMALL", 1, 1.45, 1.75, 11.80),
            ("DT4066 – DETACHER WITH LID", 1, 1.45, 1.75, 19.41),
        ],
        "OVĚŘIT",
        "Skupina DT kromě DT4002/DT4052. POZOR: kód z číselníku je jen 8místný – pro JSD "
        "doplnit na 10 míst (8505 11 10 00). DT4011 + DT4066 jsou baleny ve společném kartonu "
        "(NW 2,90 / GW 3,50 kg) – hmotnost rozdělena rovnoměrně 50/50 mezi obě položky. "
        "Ověřit antidumping na permanentní magnety z ČLR v TARIC.",
    ),
    (
        "8205 59 80 00",
        "Speciální ruční uvolňovač určený k uvolňování pevných štítků (pro etikety na brýle); "
        "uvolňovač není magnetický.",
        [("DT4002 – DT528 DETACHER NON FREQUENCY", 100, 1.20, 1.80, 107.00)],
        "",
        "Samostatný kód dle číselníku (DT4002 je výjimka ze skupiny DT).",
    ),
    (
        "7312 90 00 00",
        "Ocelová nebo poplastovaná lanka sloužící k připojení etikety ke zboží, "
        "např. provlečením.",
        [("LD3603 – LD046 LANYARD WHITE 175MM", 50_000, 86.00, 92.00, 910.00)],
        "OVĚŘIT",
        "Jediné zboží kapitoly 72/73 v zásilce → vztahuje se na něj prohlášení o neruském "
        "původu oceli (viz list Dokumenty). NENÍ zboží CBAM (příloha I nařízení 2023/956 "
        "zahrnuje 7301–7311, 7318 a 7326, nikoli 7312). Antidumping: clo na ocelová lana "
        "a kabely z ČLR se týká kódů 7312 10 81–7312 10 98 – ověřit v TARIC, že lanka "
        "skutečně patří pod 7312 90, jinak antidumping VZNIKÁ.",
    ),
    (
        "8534 00 19 00",
        "Měkký plastový štítek RF s poplachovým zabezpečením obsahující tištěný obvod, "
        "s falešným čárovým kódem, bílý nebo černý.",
        [
            ("RL4734 – RF LABEL 50*20MM 8.2MHZ", 100_000, 32.50, 35.50, 840.00),
            ("RL4620 – FREEZER RF LABEL 4*4 8.2MHZ", 300_000, 130.50, 139.50, 2250.00),
        ],
        "",
        "Sloučeny 2 položky faktury (skupina RL).",
    ),
    (
        "8531 80 70",
        "ESL – elektronické cenovky (ostatní elektrická akustická nebo vizuální "
        "signalizační zařízení).",
        [("SR98133SW – ESL 13.3\" BWR, S-LINE", 5, 3.68, 4.80, 411.75)],
        "OVĚŘIT",
        "Skupina SR98.... (kromě SR9800, SR9802DK). POZOR: kód z číselníku je jen 8místný – "
        "pro JSD doplnit na 10 míst (8531 80 70 00).",
    ),
    (
        "3926 90 97 90",
        "Ostatní výrobky z plastů – ESL držáky a kolejnice (držák na stěnu).",
        [("SR6150 – HOLDER MOUNT ON WALL", 350, 3.68, 4.40, 70.00)],
        "",
        "Skupina SR61.... dle číselníku.",
    ),
]

INCOTERM = "CFR Praha (CFR Prague)"
PALLETS_GW = 77.50
TOTAL_GW_DECLARED = 3058.00

HDR = [
    "HS kód",
    "Popis zboží (CZ)",
    "Odesílatel",
    "Země původu",
    "Množství (ks)",
    "Čistá hmotnost NW (kg)",
    "Hrubá hmotnost GW – kartony (kg)",
    "Hrubá hmotnost GW – vč. palet (kg)",
    "Celní hodnota (EUR)",
    "Incoterm",
    "Příznak",
    "Poznámka",
]

C_HDR = "FF1F3864"
C_TOT = "FFD9E1F2"
C_WARN = "FFFFE699"
C_TITLE = "FF1F3864"

thin = Side(style="thin", color="FFB0B0B0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HS Summary"

    ws.append(HDR)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=C_HDR)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER

    ctn_gw_total = sum(sum(i[3] for i in r[2]) for r in ROWS)

    # allocate pallet weight proportionally, then push the rounding remainder
    # onto the heaviest row so the column totals the declared 3 058,00 kg exactly
    gwps = [
        round(sum(i[3] for i in r[2]) * (1 + PALLETS_GW / ctn_gw_total), 2) for r in ROWS
    ]
    heaviest = max(range(len(gwps)), key=lambda i: gwps[i])
    gwps[heaviest] = round(gwps[heaviest] + TOTAL_GW_DECLARED - sum(gwps), 2)
    assert abs(sum(gwps) - TOTAL_GW_DECLARED) < 1e-9

    t_qty = t_nw = t_gw = t_gwp = t_val = 0.0
    for (hs, popis, items, flag, note), gwp in zip(ROWS, gwps):
        qty = sum(i[1] for i in items)
        nw = round(sum(i[2] for i in items), 3)
        gw = round(sum(i[3] for i in items), 3)
        val = round(sum(i[4] for i in items), 2)
        ws.append([hs, popis, SENDER, "CN", qty, nw, gw, gwp, val, INCOTERM, flag, note])
        t_qty += qty
        t_nw += nw
        t_gw += gw
        t_gwp += gwp
        t_val += val

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=True)
        row[0].font = Font(bold=True, size=10)
        for i in (4, 5, 6, 7):
            row[i].number_format = "#,##0.00"
        row[4].number_format = "#,##0"
        row[8].number_format = "#,##0.00 €"
        if row[10].value:
            row[10].fill = PatternFill("solid", fgColor=C_WARN)
            row[10].font = Font(bold=True, size=10)
            row[10].alignment = Alignment(horizontal="center", vertical="top")

    ws.append(
        [
            "CELKEM",
            f"{sum(len(r[2]) for r in ROWS)} položek faktury sloučeno do {len(ROWS)} HS kódů",
            "", "", t_qty, round(t_nw, 2), round(t_gw, 2), round(t_gwp, 2), round(t_val, 2),
            "", "", "",
        ]
    )
    tr = ws.max_row
    for c in ws[tr]:
        c.font = Font(bold=True, size=10)
        c.fill = PatternFill("solid", fgColor=C_TOT)
        c.border = BORDER
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i in (5, 6, 7, 8):
        ws.cell(tr, i).number_format = "#,##0.00"
    ws.cell(tr, 5).number_format = "#,##0"
    ws.cell(tr, 9).number_format = "#,##0.00 €"

    widths = [16, 52, 30, 10, 13, 13, 15, 15, 15, 18, 11, 62]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 46
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:L{tr}"

    # ---- legend ----
    r = tr + 2
    ws.cell(r, 1, "Legenda / kontrolní součty").font = Font(bold=True, size=11, color=C_TITLE)
    legend = [
        ("🟡 OVĚŘIT", "Kód vyžaduje kontrolu v TARIC před podáním JSD (viz sloupec Poznámka)."),
        ("CBAM", "V zásilce NENÍ žádné zboží podléhající CBAM – viz list „Dokumenty a poznámky“."),
        ("Antidumping", "Nezjištěno žádné platné antidumpingové clo; 2 kódy k ověření (7312 90 00 00, 8505 11 10)."),
        ("Sleva", "Faktura NEOBSAHUJE žádnou obchodní slevu – celní hodnota = fakturovaná hodnota."),
        ("Incoterm", f"{INCOTERM} – dle faktury (SHIP VIA: BY TRAIN). Platí pro celou zásilku."),
        ("Hrubá hmotnost", f"Kartony {ctn_gw_total:,.2f} kg + palety {PALLETS_GW:,.2f} kg = {TOTAL_GW_DECLARED:,.2f} kg "
                           f"(potvrzeno nákladním listem). Palety rozpuštěny proporcionálně dle GW kartonů."),
        ("Kontrola", f"Množství {t_qty:,.0f} ks / NW {t_nw:,.2f} kg / GW {t_gwp:,.2f} kg / "
                     f"hodnota {t_val:,.2f} EUR – vše souhlasí s fakturou i packing listem."),
    ]
    for k, v in legend:
        r += 1
        ws.cell(r, 1, k).font = Font(bold=True, size=10)
        ws.cell(r, 2, v).alignment = Alignment(wrap_text=True, vertical="top")

    # ---- sheet 2 ----
    d = wb.create_sheet("Dokumenty a poznámky")
    d.column_dimensions["A"].width = 128

    def hd(t):
        d.append([t])
        c = d.cell(d.max_row, 1)
        c.font = Font(bold=True, size=12, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=C_HDR)

    def ln(t="", bold=False):
        d.append([t])
        d.cell(d.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            d.cell(d.max_row, 1).font = Font(bold=True, size=10)

    hd("IDENTIFIKACE ZÁSILKY")
    ln("Odesílatel (sender): HANGZHOU ONTIME I.T. CO., LTD., 8th Fl, Bldg 6, Realblue Center, #18 Zixuan Rd, Hangzhou, Čína 310010")
    ln("Příjemce (CZ receiver): V-TOP EUROPE s.r.o., Františka Diviše 1012/64, 104 00 Praha 22 – Uhříněves, ČR")
    ln("IČO: 257 95 040 | DIČ/VAT: CZ25795040 | EORI: CZ25795040")
    ln("Faktura: PI24200070026043A&044&049 ze dne 17. 7. 2026 | P.O.: 260020")
    ln("Nákladní list / waybill: WPXA26B0722CZ471 | Kontejner: FCIU8451166 / plomba 319574 (40'HC)")
    ln("Přeprava: železnice (BY TRAIN), Xi'an, Čína → Praha, ČR | CFS-CFS | FREIGHT PREPAID | datum 25. 7. 2026")
    ln("Dopravce / agent: BOXLINE UCL D.O.O. / Maurice Ward & Co. s.r.o., CTP Logistic Park, Nad Kovárnou 185, 252 68 Kněževes")
    ln("Balení: 313 kartonů na 5 paletách | objem 10,77 CBM")
    ln("Původ zboží: ČÍNA (MADE IN CHINA) – uvedeno na faktuře")
    ln(f"Incoterm: {INCOTERM}")
    ln()

    hd("⚠️ PREFERENČNÍ VĚTA (COO) – NENALEZENA")
    ln("V žádném z dodaných dokumentů (faktura, packing list, statement, nákladní list) NENÍ uvedena "
       "preferenční věta o původu zboží.", bold=True)
    ln("Zboží má původ v ČÍNĚ. EU nemá s Čínou dohodu o preferenčním obchodu, preferenční původ tedy "
       "NELZE uplatnit ani doložit.")
    ln("→ Použije se standardní třetizemní (erga omnes) celní sazba dle TARIC pro každý HS kód. "
       "Preferenční zacházení se nenárokuje.")
    ln("→ Není třeba dožadovat od dodavatele – u čínského původu preferenční věta neexistuje.")
    ln()

    hd("⚠️ PROHLÁŠENÍ O NERUSKÉM PŮVODU OCELI – DOLOŽENO")
    ln("Dokument: Statement, HANGZHOU ONTIME I.T. CO., LTD., podepsáno 31. 7. 2026 (Ann Xie).", bold=True)
    ln("Odkazuje na fakturu PI24200070026043A&044&049, kontejner FCIU8451166/319574, waybill WPXA26B0722CZ471 – "
       "všechny údaje souhlasí s ostatními doklady zásilky. ✓")
    ln("Obsah: zboží kapitol HS 72 a 73 (i) nepochází z Ruska, (ii) nebylo z Ruska vyvezeno, "
       "(iii) nebylo zpracováno ze železa a oceli ruského původu (položky 7206–7229 a kapitola 73).")
    ln("Vyžadováno čl. 3g nařízení (EU) 833/2014. V zásilce se týká položky LD3603 pod kódem 7312 90 00 00.")
    ln("→ Prohlášení je platné a kompletní, přiložit k JSD.")
    ln()

    hd("CBAM – KONTROLA PODLE ČESKÉHO PŘÍJEMCE")
    ln("V zásilce NENÍ žádné zboží podléhající CBAM.", bold=True)
    ln("Jediné zboží kapitoly 72/73 je LD3603 pod kódem 7312 90 00 00. Příloha I nařízení (EU) 2023/956 "
       "zahrnuje u železa a oceli položky 7301–7311, 7318 a 7326 – kód 7312 mezi ně NEPATŘÍ. "
       "Hliník, cement, hnojiva, elektřina ani vodík se v zásilce nevyskytují.")
    ln()
    ln("Příjemce byl přesto pro pořádek dohledán v registru reference/CBAM_receivers_CZ.xlsx (list List1), řádek 43:", bold=True)
    ln("    Firma:                V-TOP EUROPE s.r.o.")
    ln("    EORI:                 CZ25795040  (souhlasí s DIČ na faktuře ✓)")
    ln("    Y KÓD:                Y137")
    ln("    Y137:                 „do 50 t ročně“ – dovoz pod limitem de minimis")
    ln("    povolení (Y128):      není uvedeno")
    ln("    podaná žádost (Y238): není uvedena")
    ln("    datum vydání:         neuvedeno")
    ln("    další reference:      25MWCCCIM008782")
    ln("→ Pokud by v budoucí zásilce CBAM zboží bylo, uvádí se kód Y137. Pro TUTO zásilku se "
       "kód CBAM do JSD neuvádí, protože žádná položka pod CBAM nespadá.")
    ln("→ Firma NEPATŘÍ mezi společnosti využívající náš sdílený účet CBAM-CZ-2025-QGM67089385721.")
    ln()

    hd("ANTIDUMPING")
    ln("Nebylo identifikováno žádné platné antidumpingové clo na položky této zásilky. "
       "Dva kódy je ale vhodné před podáním JSD ověřit v TARIC:", bold=True)
    ln("1) 7312 90 00 00 – LD3603 ocelová/poplastovaná lanka. Na ocelová lana a kabely z ČLR se "
       "antidumpingové clo vztahuje, ale pouze u kódů 7312 10 81 až 7312 10 98. Ověřit, že lanka "
       "opravdu patří pod 7312 90 00 (pomocná lanka k etiketám), nikoli pod 7312 10 – tam by "
       "antidumping vznikl.")
    ln("2) 8505 11 10 – permanentní magnety z kovu z ČLR. Ověřit aktuální stav opatření v TARIC "
       "k datu podání.")
    ln("Ostatní kódy (8531, 8534, 8205, 3926) – antidumpingová opatření na toto zboží z ČLR nejsou známa.")
    ln()

    hd("NESROVNALOSTI A BODY KE KONTROLE")
    ln("1) SCREW v nákladním listu: nákladní list WPXA26B0722CZ471 uvádí v „SAID TO CONTAIN“ mimo jiné "
       "položku SCREW (šrouby). Ve faktuře ani v packing listu ŽÁDNÁ položka šroubů není. "
       "POZOR: šrouby by spadaly pod 7318, což JE zboží CBAM. Ověřit s odesílatelem, zda jde jen o "
       "obecný popis obsahu (pin/hrot u etiket), nebo zda zásilka šrouby skutečně obsahuje.", bold=True)
    ln("2) Incoterm CFR je dle pravidel Incoterms 2020 určen pro námořní a vnitrozemskou vodní dopravu, "
       "zásilka však jede po železnici (BY TRAIN). Faktura uvádí „CFR PRAGUE“. Pro JSD použít dle faktury, "
       "případně si s dovozcem potvrdit, zda nemá být CPT Praha.")
    ln("3) DT4011 + DT4066 jsou v packing listu v jednom společném kartonu (NW 2,90 / GW 3,50 kg) bez "
       "rozpadu na jednotlivé položky. Hmotnost byla rozdělena 50/50. Obě položky spadají pod stejný "
       "HS kód 8505 11 10, na výsledek sloučeného řádku to tedy nemá vliv.")
    ln("4) Kódy 8505 11 10 a 8531 80 70 jsou v číselníku HS_kody_10.8.2026.xlsx uvedeny pouze 8místné. "
       "Pro JSD je nutné doplnit na 10 míst (8505 11 10 00, 8531 80 70 00).")
    ln("5) Faktura neobsahuje žádnou obchodní slevu – proporcionální rozpuštění slevy se neuplatňuje.")
    ln()

    hd("ZDROJE A KONTROLNÍ SOUČTY")
    ln("Zdroj HS kódů: HS_kody_10.8.2026.xlsx (list List1) – přiřazení podle prefixu kódu ONTIME "
       "(AS / DT / HD / LD / OP / RL / SR98 / SR61).")
    ln("Zdroj hmotností a množství: PACKING LIST PI24200070026043A&044&049 ze 17. 7. 2026.")
    ln("Zdroj hodnot: COMMERCIAL INVOICE PI24200070026043A&044&049 ze 17. 7. 2026.")
    ln()
    ln(f"Množství:            {t_qty:,.0f} ks     – packing list i faktura uvádějí 699 507 ks ✓")
    ln(f"Čistá hmotnost:      {t_nw:,.2f} kg   – packing list uvádí 2 792,06 kg ✓")
    ln(f"Hrubá hm. kartony:   {t_gw:,.2f} kg   – packing list uvádí 2 980,50 kg ✓")
    ln(f"Hrubá hm. s paletami:{t_gwp:,.2f} kg   – packing list i nákladní list uvádějí 3 058,00 kg ✓")
    ln(f"Celní hodnota:       {t_val:,.2f} EUR – faktura uvádí 23 787,26 EUR ✓")
    ln(f"Sloučeno:            {sum(len(r[2]) for r in ROWS)} položek faktury → {len(ROWS)} unikátních HS kódů (1 odesílatel)")

    out = "/home/user/customs/output/HS_Code_Summary_V-TOP_Europe_WPXA26B0722CZ471.xlsx"
    wb.save(out)
    print("saved:", out)
    print(f"qty={t_qty:,.0f} nw={t_nw:.2f} gw_ctn={t_gw:.2f} gw_tot={t_gwp:.2f} val={t_val:.2f}")


if __name__ == "__main__":
    main()
